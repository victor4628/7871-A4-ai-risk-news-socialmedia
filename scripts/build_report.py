"""Build the six-page editable report from verified, frozen numerical results."""
from pathlib import Path
import sys,json
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'tmp/python_deps'))
import pandas as pd
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
OUT=ROOT/'output';A=ROOT/'analysis';GREEN='254C45'
def hyperlink(p,label,url,size=None):
 rel=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel);r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),GREEN);pr.append(color)
 if size is not None:
  sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(int(size*2)));pr.append(sz)
 r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
def main():
 summary=json.loads((A/'summary.json').read_text());tests=pd.read_csv(A/'all_granger_results.csv');phase=pd.read_csv(A/'roberta_phases.csv');bias=pd.read_csv(A/'both_model_bias.csv')
 if summary.get('social_source')!='x':raise ValueError('The submission requires X data, not the legacy exploratory social source.')
 xinfo=json.loads((A/'x_selection_summary.json').read_text())
 inference=json.loads((A/'roberta_inference.json').read_text())
 audit=pd.DataFrame(json.loads((A/'both_model_audit.json').read_text())).set_index('channel')
 xreg=pd.read_csv(A/'x_selection_regressions.csv')
 primary=tests[tests.primary];significant=primary[primary.eight_primary_holm_p<.05]
 granger_claim=('Neither news nor X sentiment changes significantly Granger-predict basket returns after correcting the eight baseline tests.' if significant.empty else 'Some baseline Granger tests reject after multiplicity correction; the evidence is conditional prediction in a small retrospective sample, not structural causation.')
 vnews=pd.read_csv(A/'news_editorial_phases.csv').set_index('phase');vsocial=pd.read_csv(A/'social_phases.csv').set_index('phase')
 rb_social=phase[phase.channel=='Social'].set_index('phase')
 first,last='Sep 8-14','Sep 23-30'
 xv0,xv1=100*vsocial.loc[first,'negative_share'],100*vsocial.loc[last,'negative_share']
 xr0,xr1=100*rb_social.loc[first,'negative_share'],100*rb_social.loc[last,'negative_share']
 xtrend=('X negativity also falls under both models.' if xv1<xv0 and xr1<xr0 else 'X negativity rises under both models.' if xv1>xv0 and xr1>xr0 else 'The models disagree on the direction of X negativity.')
 d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.7);sec.right_margin=Inches(.7);sec.header_distance=Inches(.25);sec.footer_distance=Inches(.25)
 normal=d.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11);normal.font.color.rgb=RGBColor.from_string('242424');normal.paragraph_format.line_spacing=1.08;normal.paragraph_format.space_after=Pt(6)
 for name,size in [('Title',24),('Heading 1',16),('Heading 2',12)]:
  st=d.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string(GREEN);st.paragraph_format.space_after=Pt(7);st.paragraph_format.space_before=Pt(7 if name!='Title' else 0)
 d.styles['Caption'].font.name='Calibri';d.styles['Caption'].font.size=Pt(8.5);d.styles['Caption'].font.bold=False;d.styles['Caption'].font.color.rgb=RGBColor.from_string('5A6267');d.styles['Caption'].paragraph_format.space_after=Pt(7)
 # Remove decorative paragraph borders inherited from the Word template.
 for border in list(d.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 header=sec.header.paragraphs[0];header.text='AI RISK COVERAGE AND EQUITY RETURNS   |   SEPTEMBER 2026';header.style=d.styles['Normal']
 for r in header.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string(GREEN)
 footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 r=footer.add_run('October 2, 2026  •  ');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('777777')
 f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
 def p(text,style=None):
  para=d.add_paragraph(text,style)
  if style is None:para.paragraph_format.space_before=Pt(3)
  return para
 def h(text):d.add_heading(text,level=1)
 def caption(text):p(text,'Caption')
 def table(headers,rows,widths=None):
  t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
  if widths:
   for c,w in zip(t.columns,widths):c.width=Inches(w)
  for i,txt in enumerate(headers):t.rows[0].cells[i].text=str(txt)
  for row in rows:
   for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
  for ri,row in enumerate(t.rows):
   pr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');pr.append(cant)
   if ri==0:pr.append(OxmlElement('w:tblHeader'))
   for ci,c in enumerate(row.cells):
    c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if widths:c.width=Inches(widths[ci])
    cp=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),GREEN if ri==0 else ('F3F6F5' if ri%2==0 else 'FFFFFF'));cp.append(sh)
    margins=OxmlElement('w:tcMar')
    for side in ['top','bottom','left','right']:
     el=OxmlElement('w:'+side);el.set(qn('w:w'),'65' if side in ['top','bottom'] else '95');el.set(qn('w:type'),'dxa');margins.append(el)
    cp.append(margins);borders=OxmlElement('w:tcBorders')
    for side in ['top','left','bottom','right']:
     el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    cp.append(borders)
    for para in c.paragraphs:
     para.paragraph_format.space_after=Pt(0);para.paragraph_format.space_before=Pt(0);para.paragraph_format.line_spacing=1.0
     if ci>0 and len(str(row.cells[ci].text))<20:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
     for run in para.runs:run.font.size=Pt(9.3);run.font.bold=(ri==0);run.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '242424')
  return t
 def image(name,width=7.0):
  para=d.add_paragraph();para.paragraph_format.space_after=Pt(0);para.alignment=WD_ALIGN_PARAGRAPH.CENTER
  para.add_run().add_picture(str(OUT/'figures'/name),width=Inches(width))
 def page():d.add_page_break()
 # PAGE 1
 d.add_paragraph('AI Risk Coverage and Equity Returns','Title')
 byline=p('');author=byline.add_run('Victor Chen');author.font.bold=True;author.font.size=Pt(10)
 separator=byline.add_run('  |  GitHub: ');separator.font.size=Pt(9.5)
 hyperlink(byline,'github.com/victor4628/7871-A4-ai-risk-news-socialmedia','https://github.com/victor4628/7871-A4-ai-risk-news-socialmedia',size=9.5)
 p('September 8–30, 2026  |  News and X sentiment against a 23-stock AI basket')
 h('Research question and main findings')
 p(f'News coverage became less negative after the first week. {xtrend} The AI basket returned +3.72%, outperforming SPY by 4.46 percentage points and QQQ by 0.72 points, with sharply different company returns. {granger_claim} Source framing and selection bias are measurable; the sample does not reproduce the COVID study’s unusually negative US major-media pattern.')
 p('Evan Hubinger’s initiating X post gave a personal estimate above 10% for AI-caused human extinction within a decade. It was published at 21:27 EDT on September 8, after the US equity close [1]. It is an individual risk assessment, rather than an established probability or a corporate forecast. The event motivates the window; the analysis also includes ordinary AI business and technology coverage to avoid selecting only frightening stories.')
 h('Data and fixed stock universe')
 table(['Stream','Coverage after cleaning','Unit and treatment'],[
  ['News','831 headlines; 133 domains','23 daily archives; company releases excluded'],
  ['Twitter and X',f'{summary["social"]:,} posts; {summary["social_authors"]:,} authors','English; ≥100 followers and ≥100 views'],
  ['Company releases','175 headlines','Separate framing comparator'],
  ['Market prices','23 equities, SPY and QQQ','17 sessions plus September 4 baseline']], [1.12,2.15,3.83])
 caption('Sources: DataAIHub daily archives [2], TwitterAPI.io advanced search [3], and Yahoo Finance adjusted closes [4]. News and X are selected samples, not population censuses.')
 p(f'X collection uses six seeded random five-minute windows per day, capped at 100 posts each. Of {summary["raw_social"]:,} unique posts, {summary["social"]:,} remain after relevance, exposure, short-text and duplicate-author/text checks. Counts are measured at retrieval; the unfiltered sample is retained to assess selection bias.')
 table(['AI exposure','Tickers'],[
  ['Semiconductors','NVDA, AMD, AVGO, ARM, MU, TSM, ASML, MRVL'],
  ['Cloud and platforms','MSFT, GOOGL, AMZN, META, ORCL, IBM'],
  ['Enterprise software','PLTR, CRM, NOW, ADBE, SNOW'],
  ['Servers and autonomy','DELL, HPE, SMCI; TSLA']], [1.75,5.35])
 p('Each equity receives 1/23 of initial capital and is held without rebalancing. Constituents stay fixed throughout the study. The workflow follows the course’s finance case study [5], model-comparison study [6], and COVID media study [10].')
 # PAGE 2
 page();h('Sentiment trend and measurement')
 p('VADER supplies a transparent lexicon baseline [7]. Twitter RoBERTa adds a contextual model check [8]. VADER labels scores ≤−0.05 negative and ≥+0.05 positive. RoBERTa uses the highest-probability class and signed score P(positive) minus P(negative), with URL/mention placeholders. Its 2018–2021 pretraining introduces a temporal shift for 2026 text. Records receive equal weight; these are tone proxies, not direct measures of attitudes toward AI.')
 image('sentiment_models.png',6.95)
 caption('Figure 1. Daily negative-tone shares as trailing three-calendar-day averages. Solid lines use VADER and dashed lines use RoBERTa. The X panel applies the confirmed exposure filter. Unsmoothed daily values and the unfiltered comparison are supplied separately.')
 rows=[]
 vnews=pd.read_csv(A/'news_editorial_phases.csv').set_index('phase');vsocial=pd.read_csv(A/'social_phases.csv').set_index('phase')
 for phase_name in ['Sep 8-14','Sep 15-22','Sep 23-30']:
  nr=phase[(phase.phase==phase_name)&(phase.channel=='News')].iloc[0];sr=phase[(phase.phase==phase_name)&(phase.channel=='Social')].iloc[0]
  rows.append([phase_name,f'{vnews.loc[phase_name,"negative_share"]*100:.1f}%',f'{nr.negative_share*100:.1f}%',f'{vsocial.loc[phase_name,"negative_share"]*100:.1f}%',f'{sr.negative_share*100:.1f}%'])
 table(['Window','News VADER','News RoBERTa','X VADER','X RoBERTa'],rows,[1.25,1.4,1.6,1.4,1.45])
 caption('Table 1. Record-weighted negative-tone shares within three fixed calendar phases.')
 p('News negativity falls from 34.6% to 19.1% under VADER and from 17.1% to 10.7% under RoBERTa. The Hu–Liu negative-word share [9] also falls, from 4.57% to 3.27%. Existential-risk keyword coverage declines from 7.3% to 1.3% of news headlines. The result is a fading initial risk-news spike, rather than steadily worsening coverage. The final week still includes fresh safety concerns.')
 p(f'X negativity changes from {xv0:.1f}% to {xv1:.1f}% under VADER and {xr0:.1f}% to {xr1:.1f}% under RoBERTa. Mean signed X scores move from {vsocial.loc[first,"mean"]:+.3f} to {vsocial.loc[last,"mean"]:+.3f} and {rb_social.loc[first,"mean"]:+.3f} to {rb_social.loc[last,"mean"]:+.3f}, respectively. Existential-risk cues move from {100*vsocial.loc[first,"existential_share"]:.1f}% to {100*vsocial.loc[last,"existential_share"]:.1f}%. These phase comparisons describe this sample; they do not prove a response to Hubinger’s post.')
 p(f'In a fixed AI-assisted review of 20 headlines and 20 X posts, VADER/RoBERTa agreement with review labels is {100*audit.loc["News","vader_agreement"]:.0f}%/{100*audit.loc["News","roberta_agreement"]:.0f}% for news and {100*audit.loc["Social","vader_agreement"]:.0f}%/{100*audit.loc["Social","roberta_agreement"]:.0f}% for X. Labels were assigned before viewing scores, but are not independent human validation. Sarcasm, quotation and criticism of doom claims can separate negative tone from a negative stance toward AI.')
 # PAGE 3
 page();h('Equity performance and competing explanations')
 image('basket_returns.png',6.85)
 caption('Figure 2. Buy-and-hold total-return wealth, normalized at the September 4 close. That close is the prior-session baseline needed to include September 8. All results end at September 30; no October return enters the analysis.')
 p('The basket finishes +3.72%, versus SPY −0.74% and QQQ +3.00%. It reaches a −1.70% trough on September 15 and then recovers; 15 of 23 stocks finish positive. Semiconductors return +9.48% and servers +9.85%, while enterprise software returns −3.68% and cloud/platform stocks −0.26%. These category returns are equal weighted within each category.')
 image('constituent_returns_compact.png',6.9)
 caption('Figure 3. All 23 constituent total returns (%), sorted from lowest to highest across the two columns. AMD and HPE contribute +1.22 and +1.01 percentage points to the basket; Oracle contributes −0.59 points.')
 p('Oracle falls 13.53%, Salesforce 11.29% and Adobe 9.97%, while AMD rises 28.10%. The basket rises 0.61% on September 8 before the evening warning and gains 3.09% from that close to September 30. Timing rules out attributing September 8’s closing moves to the post. Oracle’s September 24 data-center execution issue [12] and the Federal Reserve’s September 16 rate increase [11] supply competing company and macro explanations. Benchmark outperformance is a performance gap, not estimated alpha or proof of a sentiment effect.')
 # PAGE 4
 page();h('Granger causality and statistical interpretation')
 p('For each channel, the restricted regression predicts daily basket log return using its own previous return. The unrestricted regression adds the previous change in the channel’s signed sentiment score. The F test asks whether this added coefficient is zero. Reverse-direction tests exchange returns and sentiment changes. One trading-day lag is the baseline; two lags are a sensitivity check [5].')
 # Native, editable linear equation, with properly subscripted symbols.
 eq=p('');eq.alignment=WD_ALIGN_PARAGRAPH.CENTER
 math=OxmlElement('m:oMath')
 def mr(text):r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);return r
 def sub(base,index):node=OxmlElement('m:sSub');e=OxmlElement('m:e');e.append(mr(base));ss=OxmlElement('m:sub');ss.append(mr(index));node.extend([e,ss]);return node
 for el in [sub('r','t'),mr(' = α + φ '),sub('r','t−1'),mr(' + β Δ'),sub('S','t−1'),mr(' + '),sub('ε','t')]:math.append(el)
 eq._p.append(math)
 p('Text and prices are aligned in New York time. News headlines lack reliable intraday timestamps, so their first archive date is assigned 23:59:59 availability. Text enters the next trading-close bin, and only the previous bin predicts a return. Weekend text pools into the next trading bin. We drop the incomplete September 8 bin and difference sentiment. This leaves 14 regression observations for each one-lag change test, with F(1,11). September 30 news and X posts after the final close cannot forecast a return inside this window.')
 table_rows=[]
 primary=tests[tests.primary]
 for model in ['VADER','RoBERTa']:
  frame=primary[primary.specification.str.startswith('RoBERTa') if model=='RoBERTa' else ~primary.specification.str.startswith('RoBERTa')]
  for _,r in frame.iterrows():
   channel='News' if 'news' in r.specification else 'X'
   direction=channel+' → return' if r.outcome=='basket_return' else 'Return → '+channel
   table_rows.append([model,direction,f'{r.F:.3f}',f'{r.p:.3f}',f'{r.eight_primary_holm_p:.3f}'])
 table(['Model','Direction','F statistic','Raw p','Holm p'],table_rows,[1.15,2.2,1.25,1.15,1.35])
 caption('Table 2. One-lag Granger tests. Holm p values correct the eight displayed tests across both models. Every regression uses the same fixed period; no lag is selected for its p value.')
 effects=[]
 for model in ['VADER','RoBERTa']:
  parts=[]
  for ch,label in [('news','news'),('social','X')]:
   r=primary[(primary.outcome=='basket_return')&primary.predictor.str.contains(ch)&(primary.specification.str.startswith('RoBERTa')==(model=='RoBERTa'))].iloc[0]
   parts.append(f'{label} {r.beta*.1:+.2f} pp [{r.ci_low*.1:+.2f}, {r.ci_high*.1:+.2f}]')
  effects.append(model+': '+', '.join(parts))
 p(granger_claim+' For a +0.10 change in signed sentiment, estimated next-return differences and 95% confidence intervals are '+ '; '.join(effects)+'. Wide intervals limit conclusions about either positive or negative impacts.')
 nominal=tests[(tests.p<.05)&~tests.primary]
 nominal_text=(f'One supplementary test has raw p<0.05' if len(nominal)==1 else f'{len(nominal)} supplementary tests have raw p<0.05')
 p(f'All {len(tests)} specifications are disclosed, including two lags, benchmark excess returns, author/conversation weighting, unfiltered X and exclusion of accounts flagged automated. {nominal_text}; the smallest Holm p across all specifications is {tests.all_specs_holm_p.min():.3f}. ADF and residual diagnostics are supplied but have little power here. Failure to reject does not establish no effect. Exposure selection uses later view/follower counts, so X tests describe retrospective prediction, not a feasible real-time trading signal.')
 # PAGE 5
 page();h('Media framing and the COVID comparison')
 p('Sacerdote, Sehgal and Cook’s 2020 working paper reports substantially more negative COVID coverage in US major media than in international major media, using word shares and classification, with controls for article length and publication week [10]. We adapt its source comparisons using headline length, archive-date fixed effects and four topic categories, with uncertainty clustered by source domain. Full article tone and audience preferences are outside this dataset.')
 image('source_models.png',6.95)
 caption('Figure 4. Negative headline shares by predeclared source group. Corporate/research releases are retained only as a comparator. US major coverage is 27.5% negative under VADER and 10.7% under RoBERTa; the international group contains only 31 headlines.')
 selection_gaps=[]
 for model in ['VADER','RoBERTa']:
  r=xreg[(xreg.model==model)&(xreg.comparison=='qualified')].iloc[0]
  selection_gaps.append(f'{100*r.coefficient:+.1f} pp [{100*r.low:.1f}, {100*r.high:.1f}]')
 table(['Comparison','VADER adjusted gap','RoBERTa adjusted gap'],[
  ['US major vs corporate','+17.1 pp [9.5, 24.6]','+7.4 pp [2.4, 12.5]'],
  ['Technology vs corporate','+13.4 pp [7.5, 19.3]','+12.1 pp [6.6, 17.7]'],
  ['US vs international major','−11.9 pp [−29.0, 5.1]','−6.2 pp [−16.7, 4.4]'],
  ['X qualified vs excluded',*selection_gaps]], [2.15,2.45,2.5])
 caption(f'Table 3. Adjusted negative-tone differences and clustered 95% confidence intervals. News: 1,006 headlines, 140 source-domain clusters. X: {xinfo["unfiltered_cleaned_posts"]:,} cleaned posts, author clusters. Both use date, topic and length controls; gaps are associations, not causal bias estimates.')
 p('The clearest measurable difference is between independent reporting and company releases. Risk/safety keywords appear in 23.9% of US-major headlines versus 2.3% of company/research headlines. The tone gap persists under both classifiers after controls, although the adjusted Hu–Liu word-share gap is imprecise. This is consistent with promotional selection and different editorial mandates. It does not show that journalists exaggerated objectively measured AI danger.')
 p('The US-versus-international result points opposite to the COVID paper and is statistically inconclusive. A negative headline may criticize doom advocates, while a reassuring denial can contain catastrophe vocabulary. These cases document framing and measurement differences; neither classifier identifies support for an extinction forecast. No objective daily AI-danger series is available to distinguish reasonable risk reporting from exaggeration.')
 p(f'X selection is measurable [13]: the filter retains {100*xinfo["retention_rate_after_common_cleaning"]:.1f}% of cleaned posts. Negative shares are {100*xinfo["vader_negative_qualified"]:.1f}% qualified versus {100*xinfo["vader_negative_unfiltered"]:.1f}% unfiltered with VADER, and {100*xinfo["roberta_negative_qualified"]:.1f}% versus {100*xinfo["roberta_negative_unfiltered"]:.1f}% with RoBERTa. Commercial-language cues appear in {100*xinfo["qualified_promotion_cue_share"]:.1f}% of qualified posts, without verifying advertisements or bots. Older posts have longer to accumulate views, introducing time-varying selection. Current exposure and platform sampling cannot establish the COVID study’s audience-demand mechanism.')
 # PAGE 6
 page();h('Conclusions and limits')
 p(f'Negative news framing retreats after the first week. {xtrend} AI equities recover in aggregate, with material differences between hardware, cloud and software companies. {granger_claim} The evidence supports source-dependent framing and measurable selection effects; it does not establish a common sentiment-driven repricing across AI companies.')
 p(f'The 17-session window leaves only 14 observations after alignment, differencing and lagging; conservative bins can miss immediate overnight reactions. X uses six random five-minute searches daily, with {xinfo["capped_windows"]} of 138 windows capped; thresholds leave {xinfo["daily_qualified_min"]}–{xinfo["daily_qualified_max"]} cleaned posts per date. Curated headlines and indexed English posts are not representative populations. RoBERTa truncates {100*inference["social"]["truncated_at_256_tokens"]/inference["social"]["n"]:.1f}% of qualified posts at 256 tokens, and review labels are AI-assisted. No extended pre-event baseline or objective daily AI-risk series supports a strong event-causal or unjustified-media-bias claim.')
 h('References and data provenance')
 refs=[
  ('[1] Hubinger, E. (2026, September 8). Initiating X post; timestamp verified through X syndication.', 'https://x.com/EvanHub/status/2097497037956891126'),
  ('[2] DataAIHub. Daily archives, September 8–30, 2026; all entries collected, with first-seen deduplication.', 'https://www.dataaihub.co/daily/2026-09-30'),
  ('[3] TwitterAPI.io. Advanced search; seeded time-stratified English X sample, verified ID boundaries and local exposure filtering.', 'https://docs.twitterapi.io/api-reference/endpoint/tweet_advanced_search'),
  ('[4] Yahoo Finance. Historical adjusted daily prices; 25 symbols, September 4–30, 2026. Endpoint checks: ChartExchange.', 'https://finance.yahoo.com/quote/NVDA/history/'),
  ('[5] Souza, T. T. P., Kolchyna, O., Treleaven, P. and Aste, T. (2015). Twitter Sentiment Analysis Applied to Finance: A Case Study in the Retail Industry.', 'https://arxiv.org/abs/1507.00784'),
  ('[6] Carvalho, J. and Plastino, A. (2020). On the evaluation and combination of state-of-the-art features in Twitter sentiment analysis.', 'https://doi.org/10.1007/s10462-020-09895-6'),
  ('[7] Hutto, C. J. and Gilbert, E. (2014). VADER: A Parsimonious Rule-Based Model for Sentiment Analysis of Social Media Text.', 'https://doi.org/10.1609/icwsm.v8i1.14550'),
  ('[8] Loureiro, D. et al. (2022). TimeLMs: Diachronic Language Models from Twitter. Cardiff NLP sentiment model; frozen quantized ONNX revision in supporting manifest.', 'https://aclanthology.org/2022.acl-demo.25/'),
  ('[9] Hu, M. and Liu, B. (2004). Mining and Summarizing Customer Reviews. Opinion lexicon distributed through NLTK.', 'https://www.cs.uic.edu/~liub/FBS/sentiment-analysis.html'),
  ('[10] Sacerdote, B., Sehgal, R. and Cook, M. (2020). Why Is All COVID-19 News Bad News? NBER Working Paper 28110; supplied November 2020 version.', 'https://www.nber.org/papers/w28110'),
  ('[11] Federal Reserve (2026, September 16). FOMC statement: a quarter-point rate increase.', 'https://www.federalreserve.gov/newsevents/pressreleases/monetary20260916a.htm'),
  ('[12] Axios (2026, September 24). Oracle moves to shield itself from AI buildout risks.', 'https://www.axios.com/2026/09/24/oracle-force-majeure-blue-owl-openai'),
  ('[13] Supporting X audit and analysis. Fixed filter policy, per-window coverage, selection comparisons, model hashes and all test specifications.', 'https://github.com/victor4628/7871-A4-ai-risk-news-socialmedia/tree/main/analysis')]
 for text,url in refs:
  para=p('');para.paragraph_format.space_after=Pt(4);para.paragraph_format.line_spacing=1.0;rr=para.add_run(text+' ');rr.font.size=Pt(9);hyperlink(para,'Source',url,size=9)
 p(f'News and prices were collected October 1, 2026; X was retrieved October 2. The support archive supplies IDs, URLs, metadata, daily measures, all {len(tests)} test specifications and reproducible scripts. Full post bodies, scraped HTML, API credentials and model weights are omitted from distribution.','Caption')
 d.core_properties.title='AI Risk Coverage and Equity Returns';d.core_properties.subject='September 8–30, 2026';d.core_properties.author='Victor Chen';d.core_properties.keywords='AI sentiment, Granger causality, media framing'
 OUT.mkdir(exist_ok=True);path=OUT/'AI_Risk_Sentiment_2026_Report.docx';d.save(path)
 print('SAVED',path)
if __name__=='__main__':main()
